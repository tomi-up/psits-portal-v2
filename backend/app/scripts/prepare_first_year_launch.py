"""Prepare first-year staging data for launch; dry-run unless --commit is used.

Uses PRODUCTION_DATABASE_URL when configured, otherwise the first-year sheet
in private_data/psits_attendance_with_emails.xlsx. An optional CSV/XLSX can
override emails. It also creates two PHP 100 fees and can exempt Year 1 from
one past event.
"""

import argparse
import csv
import re
import uuid
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.balance import MembershipFee
from app.models.event import Event
from app.models.student import Program, SchoolYear, Student, StudentSchoolYear

PROJECT_ROOT = Path(__file__).resolve().parents[3]
ATTENDANCE_WORKBOOK = PROJECT_ROOT / "private_data" / "psits_attendance_with_emails.xlsx"
SCHOOL_YEAR_LABEL = "2026-2027"
FEE_AMOUNT = Decimal("100.00")
SEMESTERS = ("1ST", "2ND")
PROGRAM_NAMES = {
    "BSCS": "Bachelor of Science in Computer Science",
    "BSIT": "Bachelor of Science in Information Technology",
    "BLIS": "Bachelor of Library and Information Science",
    "BSINFOSYS": "Bachelor of Science in Information Systems",
}


def _header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def load_email_map(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    if not path.exists():
        raise ValueError(f"Email file not found: {path}")

    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.reader(handle))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
    else:
        raise ValueError("Email file must be CSV or XLSX")

    header_row = id_column = email_column = None
    for index, row in enumerate(rows[:30]):
        names = [_header(value) for value in row]
        id_column = next(
            (i for i, name in enumerate(names) if name in {"studentid", "studentnumber", "idnumber"}),
            None,
        )
        email_column = next(
            (i for i, name in enumerate(names) if name in {"email", "emailaddress", "studentemail"}),
            None,
        )
        if id_column is not None and email_column is not None:
            header_row = index
            break
    if header_row is None or id_column is None or email_column is None:
        raise ValueError("Could not find Student ID and Email columns in the email file")

    result: dict[str, str] = {}
    for row in rows[header_row + 1 :]:
        if max(id_column, email_column) >= len(row):
            continue
        student_id = str(row[id_column] or "").strip().upper()
        email = str(row[email_column] or "").strip().lower()
        if not student_id and not email:
            continue
        if not student_id or "@" not in email:
            raise ValueError("The email file contains an invalid Student ID or email")
        if student_id in result and result[student_id] != email:
            raise ValueError(f"The email file has conflicting emails for {student_id}")
        result[student_id] = email
    return result


def load_attendance_workbook() -> tuple[list[dict[str, object]], int, int]:
    if not ATTENDANCE_WORKBOOK.exists():
        raise ValueError(f"Attendance workbook not found: {ATTENDANCE_WORKBOOK}")

    workbook = load_workbook(ATTENDANCE_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["1ST YR"]
    records: list[dict[str, object]] = []
    skipped = 0
    missing_emails = 0
    seen_ids: set[str] = set()
    seen_emails: set[str] = set()
    for row in sheet.iter_rows(min_row=8, values_only=True):
        student_id = str(row[1] or "").strip().upper()
        if not re.fullmatch(r"\d{2}-\d{5}", student_id):
            continue
        values = {
            "student_id": student_id,
            "last_name": str(row[2] or "").strip().upper(),
            "first_name": str(row[3] or "").strip().upper(),
            "middle_name": str(row[4] or "").strip().upper() or None,
            "year": str(row[5] or "").strip().upper(),
            "course": str(row[6] or "").strip().upper(),
            "section": str(row[7] or "").strip().upper(),
            "email": str(row[8] or "").strip().lower(),
        }
        required = ("last_name", "first_name", "year", "course", "section")
        if any(not values[field] for field in required) or values["year"] != "1ST":
            skipped += 1
            continue
        email = str(values["email"])
        if email and not email.endswith("@usm.edu.ph"):
            raise ValueError("A first-year email is outside the required @usm.edu.ph domain")
        if student_id in seen_ids or (email and email in seen_emails):
            raise ValueError("The first-year workbook contains a duplicate Student ID or email")
        seen_ids.add(student_id)
        if email:
            seen_emails.add(email)
        else:
            missing_emails += 1
            values["email"] = None
        records.append({key: value for key, value in values.items() if key != "year"})
    return records, skipped, missing_emails


def load_source() -> tuple[list[dict[str, object]], str, int, int]:
    source_url = (settings.production_database_url or "").strip()
    if not source_url:
        rows, skipped, missing_emails = load_attendance_workbook()
        return rows, "attendance workbook with emails", skipped, missing_emails
    if source_url == settings.database_url:
        raise ValueError("PRODUCTION_DATABASE_URL and DATABASE_URL must not be the same")

    source_engine = create_engine(source_url, pool_pre_ping=True)
    with Session(source_engine) as source:
        records = (
            source.query(Student, StudentSchoolYear, Program)
            .join(StudentSchoolYear, StudentSchoolYear.student_id == Student.id)
            .join(SchoolYear, SchoolYear.id == StudentSchoolYear.school_year_id)
            .join(Program, Program.id == StudentSchoolYear.program_id)
            .filter(SchoolYear.label == SCHOOL_YEAR_LABEL, StudentSchoolYear.year_level == 1)
            .all()
        )
        rows = [
            {
                "student_id": student.student_id,
                "first_name": student.first_name,
                "middle_name": student.middle_name,
                "last_name": student.last_name,
                "email": student.email,
                "course": program.code.upper(),
                "section": enrollment.section,
            }
            for student, enrollment, program in records
        ]
    source_engine.dispose()
    return rows, "production database", 0, sum(not row.get("email") for row in rows)


def prepare(commit: bool, emails_path: Path | None, event_id: str | None) -> None:
    if settings.environment.lower() != "staging":
        raise ValueError("Refusing to run: ENVIRONMENT must be staging")

    source_rows, source_label, skipped_source_rows, missing_source_emails = load_source()
    email_map = load_email_map(emails_path)
    source_ids = {str(row["student_id"]) for row in source_rows}
    unmatched_emails = set(email_map) - source_ids
    if unmatched_emails:
        raise ValueError(f"Email file has {len(unmatched_emails)} IDs absent from the first-year source")

    counts = {key: 0 for key in (
        "students_created", "students_updated", "enrollments_created",
        "enrollments_updated", "emails_applied", "fees_created", "fees_updated",
        "programs_created",
    )}
    target_engine = create_engine(settings.database_url, pool_pre_ping=True)
    with Session(target_engine) as db:
        try:
            school_year = db.query(SchoolYear).filter(SchoolYear.label == SCHOOL_YEAR_LABEL).first()
            if not school_year:
                raise ValueError(f"School year {SCHOOL_YEAR_LABEL} does not exist in staging")
            programs = {program.code.upper(): program for program in db.query(Program).all()}
            missing_programs = {str(row["course"]) for row in source_rows} - set(programs)
            unknown_programs = missing_programs - set(PROGRAM_NAMES)
            if unknown_programs:
                raise ValueError(f"Unknown source programs: {sorted(unknown_programs)}")
            for code in sorted(missing_programs):
                program = Program(id=str(uuid.uuid4()), code=code, name=PROGRAM_NAMES[code])
                db.add(program)
                programs[code] = program
                counts["programs_created"] += 1

            event = db.query(Event).filter(Event.id == event_id).first() if event_id else None
            if event_id and event is None:
                raise ValueError("The selected exemption event does not exist in staging")

            students = {
                student.student_id: student
                for student in db.query(Student).filter(Student.student_id.in_(source_ids)).all()
            }
            existing_database_ids = [student.id for student in students.values()]
            enrollments = {
                enrollment.student_id: enrollment
                for enrollment in db.query(StudentSchoolYear)
                .filter(
                    StudentSchoolYear.school_year_id == school_year.id,
                    StudentSchoolYear.student_id.in_(existing_database_ids),
                )
                .all()
            }
            fees = {
                (fee.student_id, fee.semester): fee
                for fee in db.query(MembershipFee)
                .filter(
                    MembershipFee.school_year_id == school_year.id,
                    MembershipFee.student_id.in_(existing_database_ids),
                )
                .all()
            }
            for row in source_rows:
                student_id = str(row["student_id"])
                email = email_map.get(student_id) or row.get("email")
                student = students.get(student_id)
                if student is None:
                    student = Student(
                        id=str(uuid.uuid4()), student_id=student_id,
                        first_name=str(row["first_name"]), middle_name=row.get("middle_name"),
                        last_name=str(row["last_name"]), email=email, is_active=False,
                    )
                    db.add(student)
                    students[student_id] = student
                    counts["students_created"] += 1
                    counts["emails_applied"] += int(bool(email))
                else:
                    changed = False
                    for field in ("first_name", "middle_name", "last_name"):
                        if getattr(student, field) != row.get(field):
                            setattr(student, field, row.get(field))
                            changed = True
                    counts["students_updated"] += int(changed)
                if email and student.email != email:
                    student.email = str(email)
                    counts["emails_applied"] += 1

                enrollment = enrollments.get(student.id)
                program = programs[str(row["course"])]
                desired = (program.id, 1, row.get("section"), "ACTIVE", "REGULAR")
                if enrollment is None:
                    enrollment = StudentSchoolYear(
                        id=str(uuid.uuid4()), student_id=student.id, school_year_id=school_year.id,
                        program_id=desired[0], year_level=1, section=desired[2],
                        status="ACTIVE", academic_standing="REGULAR",
                    )
                    db.add(enrollment)
                    enrollments[student.id] = enrollment
                    counts["enrollments_created"] += 1
                else:
                    current = (enrollment.program_id, enrollment.year_level, enrollment.section,
                               enrollment.status, enrollment.academic_standing)
                    if current != desired:
                        (enrollment.program_id, enrollment.year_level, enrollment.section,
                         enrollment.status, enrollment.academic_standing) = desired
                        counts["enrollments_updated"] += 1

                for semester in SEMESTERS:
                    fee = fees.get((student.id, semester))
                    if fee is None:
                        fee = MembershipFee(
                            id=str(uuid.uuid4()), student_id=student.id,
                            school_year_id=school_year.id, semester=semester,
                            amount_due=FEE_AMOUNT, amount_paid=Decimal("0.00"),
                        )
                        db.add(fee)
                        fees[(student.id, semester)] = fee
                        counts["fees_created"] += 1
                    elif Decimal(fee.amount_paid) > FEE_AMOUNT:
                        raise ValueError(f"A {semester} fee has payments above the PHP 100 due")
                    elif Decimal(fee.amount_due) != FEE_AMOUNT:
                        fee.amount_due = FEE_AMOUNT
                        counts["fees_updated"] += 1

            if event:
                event.excused_year_levels = sorted(set(event.excused_year_levels or []) | {1})
            db.flush()
            db.commit() if commit else db.rollback()
        except Exception:
            db.rollback()
            raise
    target_engine.dispose()

    print(f"Mode: {'COMMITTED' if commit else 'DRY RUN (rolled back)'}")
    print(f"Source: {source_label}")
    print(f"First-year source records: {len(source_rows)}")
    print(f"Incomplete source records skipped: {skipped_source_rows}")
    print(f"Source records missing email: {missing_source_emails}")
    for key, value in counts.items():
        print(f"{key}: {value}")
    print(f"event_year_1_exemption: {'yes' if event_id else 'not requested'}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--emails", type=Path, help="CSV/XLSX with Student ID and Email columns")
    parser.add_argument("--event-id", help="Past event that did not require first-year attendance")
    parser.add_argument("--commit", action="store_true", help="Commit changes; default is rollback")
    args = parser.parse_args()
    prepare(args.commit, args.emails, args.event_id)
