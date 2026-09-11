"""Builds the Excel (.xlsx) survey results report for an event.

Renders exactly the aggregates produced by
app.api.v1.endpoints.admin_events.get_survey_results - never queries the
database itself, so the on-screen results and the export can't drift apart.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from app.api.v1.endpoints.admin_events import SurveyResultsResponse

ORG_NAME = "PHILIPPINE SOCIETY OF INFORMATION TECHNOLOGY STUDENTS"
LOGO_PATH = Path(__file__).resolve().parents[3] / "frontend" / "public" / "psits-logo.png"

THIN_BORDER = Border(*(Side(style="thin", color="B0B7C3") for _ in range(4)))
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")

SUMMARY_COLUMNS = [
    ("Section", 40),
    ("Statement", 65),
    ("Average Rating", 16),
    ("Responses", 12),
]

COMMENT_COLUMNS = [
    ("Student ID", 16),
    ("Student Name", 28),
    ("Comment", 70),
    ("Submitted At", 20),
]


def _write_header(ws: Worksheet, event_name: str, subtitle: str, last_col: str) -> int:
    row = 1

    if LOGO_PATH.exists():
        img = XLImage(str(LOGO_PATH))
        img.width = 64
        img.height = 64
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 48
        ws.row_dimensions[2].height = 48
        title_col = "B"
    else:
        title_col = "A"

    ws.merge_cells(f"{title_col}1:{last_col}1")
    ws[f"{title_col}1"] = ORG_NAME
    ws[f"{title_col}1"].font = Font(bold=True, size=13)

    ws.merge_cells(f"{title_col}2:{last_col}2")
    ws[f"{title_col}2"] = subtitle
    ws[f"{title_col}2"].font = Font(bold=True, size=11, color="475569")

    row = 4
    ws.merge_cells(f"A{row}:{last_col}{row}")
    ws[f"A{row}"] = event_name
    ws[f"A{row}"].font = Font(bold=True, size=14)

    return row + 2


def build_survey_results_workbook(data: "SurveyResultsResponse") -> BytesIO:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    last_col = get_column_letter(len(SUMMARY_COLUMNS))

    next_row = _write_header(summary_ws, data.event_name, "SURVEY RESULTS SUMMARY", last_col)

    summary_ws[f"A{next_row}"] = "Total Responses"
    summary_ws[f"A{next_row}"].font = Font(bold=True)
    summary_ws[f"B{next_row}"] = f"{data.total_responses} / {data.total_eligible}"
    summary_ws[f"B{next_row}"].font = Font(bold=True)
    next_row += 2

    header_row = next_row
    for col_idx, (title, width) in enumerate(SUMMARY_COLUMNS, start=1):
        cell = summary_ws.cell(row=header_row, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        summary_ws.column_dimensions[get_column_letter(col_idx)].width = width

    row = header_row + 1
    for section in data.sections:
        for q in section.questions:
            values = [
                section.title,
                q.text,
                q.average if q.average is not None else "—",
                q.responses,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = summary_ws.cell(row=row, column=col_idx, value=value)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(
                    horizontal="left" if col_idx in (1, 2) else "center",
                    vertical="center",
                    wrap_text=col_idx == 2,
                )
            row += 1

    summary_ws.auto_filter.ref = f"A{header_row}:{last_col}{row - 1}"
    summary_ws.freeze_panes = f"A{header_row + 1}"

    comments_ws = wb.create_sheet("Comments")
    comments_last_col = get_column_letter(len(COMMENT_COLUMNS))
    c_row = _write_header(comments_ws, data.event_name, "SURVEY COMMENTS", comments_last_col)
    c_row += 1

    header_row = c_row
    for col_idx, (title, width) in enumerate(COMMENT_COLUMNS, start=1):
        cell = comments_ws.cell(row=header_row, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        comments_ws.column_dimensions[get_column_letter(col_idx)].width = width

    row = header_row + 1
    for c in data.comments:
        values = [c.student_id, c.student_name, c.comment, c.submitted_at.strftime("%B %d, %Y %I:%M %p")]
        for col_idx, value in enumerate(values, start=1):
            cell = comments_ws.cell(row=row, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="left" if col_idx in (2, 3) else "center", vertical="center", wrap_text=col_idx == 3
            )
        row += 1

    if data.comments:
        comments_ws.auto_filter.ref = f"A{header_row}:{comments_last_col}{row - 1}"
    comments_ws.freeze_panes = f"A{header_row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
